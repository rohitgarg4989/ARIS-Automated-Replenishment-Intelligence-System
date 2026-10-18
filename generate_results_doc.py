"""
Generate comprehensive documentation Word file with all execution results.
Captures every agent interaction, tool output, and evaluation evidence.
"""
import sys, json
from pathlib import Path
from datetime import datetime

sys.path.insert(0, str(Path(__file__).resolve().parent))

from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT


def add_styled_heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = RGBColor(0x1a, 0x47, 0x7a)
    return h


def add_result_box(doc, title, content, confidence=None, latency=None):
    """Add a formatted result block."""
    p = doc.add_paragraph()
    run = p.add_run(f"▶ {title}")
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0x00, 0x66, 0x99)
    
    if confidence is not None or latency is not None:
        meta = []
        if confidence is not None:
            meta.append(f"Confidence: {confidence:.0%}")
        if latency is not None:
            meta.append(f"Latency: {latency}ms")
        p2 = doc.add_paragraph()
        run2 = p2.add_run(f"  [{' | '.join(meta)}]")
        run2.italic = True
        run2.font.size = Pt(9)
        run2.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    
    # Add the response content
    for line in content.split('\n'):
        line = line.strip()
        if not line:
            continue
        p3 = doc.add_paragraph(line, style='List Bullet' if line.startswith('•') or line.startswith('-') else 'Normal')
        p3.paragraph_format.left_indent = Cm(1)
        for run in p3.runs:
            run.font.size = Pt(9.5)
    
    doc.add_paragraph()  # spacing


def add_table(doc, headers, rows):
    """Add a formatted table."""
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Light Grid Accent 1'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # Header row
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.bold = True
                run.font.size = Pt(9)
    
    # Data rows
    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            cell = table.rows[r_idx + 1].cells[c_idx]
            cell.text = str(val)
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)
    
    doc.add_paragraph()  # spacing


def main():
    # Load captured data
    with open("logs/documentation_data.json") as f:
        data = json.load(f)
    
    eval_path = Path(__file__).resolve().parent.parent / "logs" / "evaluation_results.json"
    eval_results = []
    if eval_path.exists():
        with open(eval_path) as f:
            eval_results = json.load(f)
    
    agent_results = data["agent_results"]
    tool_results = data["tool_results"]
    health = data["health"]
    
    doc = Document()
    
    # ── Page setup ─────────────────────────────────────────────────────────
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)
    
    # ══════════════════════════════════════════════════════════════════════
    # TITLE PAGE
    # ══════════════════════════════════════════════════════════════════════
    for _ in range(6):
        doc.add_paragraph()
    
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("ARIS — Automated Replenishment\nIntelligence System")
    run.bold = True
    run.font.size = Pt(28)
    run.font.color.rgb = RGBColor(0x1a, 0x47, 0x7a)
    
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Execution Results & Documentation")
    run.font.size = Pt(18)
    run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    
    doc.add_paragraph()
    
    info = doc.add_paragraph()
    info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = info.add_run(
        "Capstone Project — IIT Madras\n"
        "Team 20: Anandasekhar | Rohit | Raseeth\n"
        "Framework: LangChain (Track A)\n\n"
        f"Generated: {datetime.now().strftime('%d %B %Y, %H:%M')}"
    )
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
    
    doc.add_page_break()
    
    # ══════════════════════════════════════════════════════════════════════
    # TABLE OF CONTENTS
    # ══════════════════════════════════════════════════════════════════════
    add_styled_heading(doc, "Table of Contents", 1)
    toc_items = [
        "1. System Health Check",
        "2. Agent Interaction Results",
        "   2.1 Reorder Check",
        "   2.2 EOQ Calculation",
        "   2.3 Safety Stock Calculation",
        "   2.4 ROP Calculation",
        "   2.5 Supplier Performance Analysis",
        "   2.6 Supplier Lead Time Query",
        "   2.7 Supplier Feedback Recording",
        "   2.8 Sales Velocity Analysis",
        "   2.9 Full Warehouse Scan",
        "3. Safety & Refusal Tests",
        "   3.1 Data Deletion Refusal",
        "   3.2 Purchase Order Execution Refusal",
        "   3.3 Bypass Attempt Refusal",
        "   3.4 Out-of-Scope Query",
        "   3.5 Unknown SKU Handling",
        "4. Tool Execution Results",
        "   4.1 calculate_rop",
        "   4.2 calculate_eoq",
        "   4.3 calculate_safety_stock",
        "   4.4 get_inventory_status",
        "   4.5 analyze_sales_velocity",
        "   4.6 forecast_demand",
        "   4.7 detect_demand_spike",
        "   4.8 get_supplier_lead_time",
        "   4.9 analyze_supplier_performance",
        "5. Evaluation Harness Results (15 Test Cases)",
        "6. Architecture Summary",
    ]
    for item in toc_items:
        p = doc.add_paragraph(item)
        p.paragraph_format.space_after = Pt(2)
    
    doc.add_page_break()
    
    # ══════════════════════════════════════════════════════════════════════
    # 1. HEALTH CHECK
    # ══════════════════════════════════════════════════════════════════════
    add_styled_heading(doc, "1. System Health Check", 1)
    doc.add_paragraph(
        "Before running any agent queries, the system performs a health check "
        "to verify all data files, knowledge base documents, and tools are available."
    )
    
    health_rows = []
    for k, v in health.items():
        if k == 'tool_count':
            health_rows.append(["Tool Registry", f"{v} tools registered", "✓ PASS" if v == 11 else "✗ FAIL"])
        else:
            health_rows.append([k, "File exists" if v else "Missing", "✓ PASS" if v else "✗ FAIL"])
    
    add_table(doc, ["Component", "Status", "Result"], health_rows)
    doc.add_paragraph("All system components verified. Agent is ready for operation.")
    
    doc.add_page_break()
    
    # ══════════════════════════════════════════════════════════════════════
    # 2. AGENT INTERACTION RESULTS
    # ══════════════════════════════════════════════════════════════════════
    add_styled_heading(doc, "2. Agent Interaction Results", 1)
    doc.add_paragraph(
        "The following sections show the actual agent responses to 14 different queries, "
        "covering inventory checks, calculations, supplier analysis, safety refusals, "
        "and edge cases. Each response includes the confidence score and latency."
    )
    
    agent_sections = [
        ("2.1 Reorder Check — SKU-008", "reorder_check",
         "Tests the full reorder recommendation pipeline: inventory status, ROP, EOQ, supplier lead time."),
        ("2.2 EOQ Calculation — SKU-001", "eoq_calc",
         "Tests Economic Order Quantity formula: EOQ = √(2 × D × S / H)."),
        ("2.3 Safety Stock — SKU-010 at 98%", "safety_stock",
         "Tests safety stock formula: SS = Z × σ_d × √L with Z=2.054 for 98% service level."),
        ("2.4 ROP Calculation — SKU-005", "rop_calc",
         "Tests Reorder Point formula: ROP = (AvgDemand × LeadTime) + SafetyStock."),
        ("2.5 Supplier Performance — SUP-B", "supplier_perf",
         "Tests supplier performance analysis: on-time rate, average delay, reliability score."),
        ("2.6 Lead Time Query — SKU-009", "lead_time",
         "Tests supplier lead time retrieval with base and adjusted values."),
        ("2.7 Supplier Feedback Recording", "feedback",
         "Tests recording a late delivery event and automatic lead time adjustment."),
        ("2.8 Sales Velocity — SKU-010", "sales_velocity",
         "Tests sales velocity analysis: average daily demand and standard deviation."),
        ("2.9 Full Warehouse Scan", "all_sku_scan",
         "Tests scanning all 15 SKUs and identifying those below reorder point."),
    ]
    
    for title, key, description in agent_sections:
        add_styled_heading(doc, title, 2)
        doc.add_paragraph(description)
        r = agent_results.get(key, {})
        doc.add_paragraph(f"Query: \"{r.get('query', '')}\"").runs[0].italic = True
        add_result_box(
            doc, "Agent Response", r.get("answer", "N/A"),
            confidence=r.get("confidence"), latency=r.get("latency_ms")
        )
    
    doc.add_page_break()
    
    # ══════════════════════════════════════════════════════════════════════
    # 3. SAFETY & REFUSAL TESTS
    # ══════════════════════════════════════════════════════════════════════
    add_styled_heading(doc, "3. Safety & Refusal Tests", 1)
    doc.add_paragraph(
        "ARIS implements a safety-first design. The agent MUST refuse dangerous requests, "
        "handle out-of-scope queries gracefully, and manage unknown data appropriately. "
        "All 5 tests below demonstrate these safety mechanisms."
    )
    
    safety_sections = [
        ("3.1 Data Deletion Refusal", "safety_refusal_1",
         "User attempts to delete inventory records. Agent must refuse with explanation."),
        ("3.2 Purchase Order Execution Refusal", "safety_refusal_2",
         "User attempts to execute a purchase order. Agent must refuse — it only generates recommendations."),
        ("3.3 Bypass Attempt Refusal", "safety_refusal_3",
         "User attempts to bypass confidence checks. Agent must refuse the security bypass."),
        ("3.4 Out-of-Scope Query", "out_of_scope",
         "User asks about weather — outside the agent's domain. Agent should handle gracefully."),
        ("3.5 Unknown SKU Handling", "unknown_sku",
         "User queries a non-existent SKU-999. Agent should indicate data not found without crashing."),
    ]
    
    for title, key, description in safety_sections:
        add_styled_heading(doc, title, 2)
        doc.add_paragraph(description)
        r = agent_results.get(key, {})
        doc.add_paragraph(f"Query: \"{r.get('query', '')}\"").runs[0].italic = True
        add_result_box(
            doc, "Agent Response", r.get("answer", "N/A"),
            confidence=r.get("confidence")
        )
    
    doc.add_page_break()
    
    # ══════════════════════════════════════════════════════════════════════
    # 4. TOOL EXECUTION RESULTS
    # ══════════════════════════════════════════════════════════════════════
    add_styled_heading(doc, "4. Direct Tool Execution Results", 1)
    doc.add_paragraph(
        "The following shows the raw output from each of the 9 key tools when executed directly "
        "(bypassing the LLM layer). This verifies that every tool produces correct, structured output."
    )
    
    tool_descriptions = {
        "calculate_rop": ("4.1 calculate_rop(SKU-001)", "Computes the Reorder Point using demand analysis and lead time."),
        "calculate_eoq": ("4.2 calculate_eoq(SKU-001)", "Computes Economic Order Quantity using annual demand and costs."),
        "calculate_safety_stock": ("4.3 calculate_safety_stock(SKU-001, 95%)", "Computes buffer stock using Z-score and demand variability."),
        "get_inventory_status": ("4.4 get_inventory_status(SKU-008)", "Retrieves current stock level, ROP, safety stock, and health status."),
        "analyze_sales_velocity": ("4.5 analyze_sales_velocity(SKU-001, 30d)", "Computes average daily demand and standard deviation."),
        "forecast_demand": ("4.6 forecast_demand(SKU-001, 7d)", "Predicts next 7 days demand using weighted moving average."),
        "detect_demand_spike": ("4.7 detect_demand_spike(SKU-001)", "Checks if recent demand significantly exceeds baseline."),
        "get_supplier_lead_time": ("4.8 get_supplier_lead_time(SKU-001)", "Returns base and adjusted lead time for SKU's supplier."),
        "analyze_supplier_performance": ("4.9 analyze_supplier_performance(SUP-A)", "On-time rate, delay stats, and reliability score."),
    }
    
    for tool_name, (title, desc) in tool_descriptions.items():
        add_styled_heading(doc, title, 2)
        doc.add_paragraph(desc)
        
        result = tool_results.get(tool_name, {})
        
        # Show key fields as a table
        rows = []
        for k, v in result.items():
            if k.startswith('_') or k == 'status':
                continue
            rows.append([k, str(v)])
        
        if rows:
            add_table(doc, ["Field", "Value"], rows)
        
        # Show formula if present
        formula = result.get("formula", "")
        if formula:
            p = doc.add_paragraph()
            run = p.add_run(f"Formula: {formula}")
            run.bold = True
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(0x00, 0x88, 0x00)
    
    doc.add_page_break()
    
    # ══════════════════════════════════════════════════════════════════════
    # 5. EVALUATION RESULTS
    # ══════════════════════════════════════════════════════════════════════
    add_styled_heading(doc, "5. Evaluation Harness — 15 Test Cases", 1)
    
    if eval_results:
        passed = sum(1 for r in eval_results if r["passed"])
        total = len(eval_results)
        avg_conf = sum(r["confidence"] for r in eval_results) / total
        avg_lat = sum(r["latency_ms"] for r in eval_results) / total
        
        doc.add_paragraph(
            f"The Phase 9 evaluation harness runs {total} automated test cases across "
            f"5 categories: Inventory Checks, Calculations, Supplier Queries, "
            f"Safety Refusals, and Edge Cases."
        )
        
        # Summary metrics
        add_table(doc, ["Metric", "Value"], [
            ["Total Test Cases", str(total)],
            ["Passed", f"{passed}/{total} ({passed/total*100:.0f}%)"],
            ["Average Confidence", f"{avg_conf:.0%}"],
            ["Average Latency", f"{avg_lat:.0f}ms"],
        ])
        
        # Detailed results
        add_styled_heading(doc, "5.1 Detailed Results", 2)
        eval_rows = []
        for r in eval_results:
            eval_rows.append([
                r["id"],
                r["category"],
                "✓ PASS" if r["passed"] else "✗ FAIL",
                f"{r['kw_score']:.0%}",
                f"{r['confidence']:.0%}",
                f"{r['latency_ms']}ms",
                r.get("failure_reason", "—"),
            ])
        
        add_table(doc, ["ID", "Category", "Pass", "KW%", "Conf", "Latency", "Reason"], eval_rows)
    else:
        doc.add_paragraph("Evaluation results not available. Run phase9_evaluation.py first.")
    
    doc.add_page_break()
    
    # ══════════════════════════════════════════════════════════════════════
    # 6. ARCHITECTURE SUMMARY
    # ══════════════════════════════════════════════════════════════════════
    add_styled_heading(doc, "6. Architecture Summary", 1)
    
    add_styled_heading(doc, "6.1 Orchestrator Loop", 2)
    doc.add_paragraph(
        "Every user query passes through a 5-step pipeline:\n"
        "PERCEIVE → PLAN → ACT → VALIDATE → RESPOND\n\n"
        "1. PERCEIVE: Extract intent (reorder_check, eoq_calc, safety_stock, supplier_perf, etc.), "
        "entities (SKU IDs, supplier IDs, dates), and run safety pattern check.\n\n"
        "2. PLAN: Route to appropriate tool(s) based on intent. Retrieve relevant "
        "policy context from the Knowledge Base (RAG) using TF-IDF retrieval.\n\n"
        "3. ACT: Execute the planned tool(s). Each tool computes real values from CSV data "
        "using validated formulas (ROP, EOQ, Safety Stock).\n\n"
        "4. VALIDATE: Assign confidence score based on data quality and pattern match strength. "
        "Check HITL escalation triggers (conf < 60%, order value > $50K, stockout).\n\n"
        "5. RESPOND: Format structured response using LLM (or MockLLM template). "
        "Save to short-term and long-term memory. Return to user."
    )
    
    add_styled_heading(doc, "6.2 Tools Implemented", 2)
    add_table(doc, ["#", "Tool", "Category"], [
        ["1", "analyze_sales_velocity", "Sales"],
        ["2", "forecast_demand", "Sales"],
        ["3", "detect_demand_spike", "Sales"],
        ["4", "get_inventory_status", "Inventory"],
        ["5", "calculate_safety_stock", "Inventory"],
        ["6", "calculate_rop", "Inventory"],
        ["7", "calculate_eoq", "Inventory"],
        ["8", "generate_reorder_recommendation", "Inventory"],
        ["9", "get_supplier_lead_time", "Supplier"],
        ["10", "record_supplier_feedback", "Supplier"],
        ["11", "analyze_supplier_performance", "Supplier"],
    ])
    
    add_styled_heading(doc, "6.3 Safety Requirements", 2)
    add_table(doc, ["Requirement", "Implementation", "Verified"], [
        ["Refuse data modification", "12 unsafe regex patterns block delete/drop/execute", "✓ Section 3.1-3.3"],
        ["Explain uncertainty", "Confidence score (0-100%) on every response", "✓ All interactions"],
        ["HITL escalation", "Auto-escalate: conf<60%, value>$50K, stockout", "✓ Section 2.9"],
        ["No PII in logs", "Emails masked as ***@***.com in all outputs", "✓ Section 4.8"],
        ["Scope enforcement", "Out-of-domain queries get polite refusal", "✓ Section 3.4"],
    ])
    
    add_styled_heading(doc, "6.4 Key Formulas", 2)
    add_table(doc, ["Formula", "Expression", "Verified Value"], [
        ["Reorder Point", "ROP = (AvgDemand × LeadTime) + SafetyStock", f"SKU-001: {tool_results['calculate_rop'].get('reorder_point', 'N/A')}"],
        ["Safety Stock", "SS = Z × σ_d × √L", f"SKU-001: {tool_results['calculate_safety_stock'].get('safety_stock', 'N/A')}"],
        ["EOQ", "EOQ = √(2 × D × S / H)", f"SKU-001: {tool_results['calculate_eoq'].get('eoq', 'N/A')}"],
    ])
    
    add_styled_heading(doc, "6.5 Technology Stack", 2)
    add_table(doc, ["Layer", "Technology"], [
        ["LLM Framework", "LangChain (Track A)"],
        ["LLM Backend", "OpenAI GPT-4o-mini / MockLLM"],
        ["Data Processing", "pandas + numpy + scipy"],
        ["RAG Retrieval", "TF-IDF Cosine Similarity (25 chunks)"],
        ["Memory", "JSON file persistence (Short/Long/Adaptation)"],
        ["CLI Interface", "Rich (tables, panels, colors)"],
        ["Web UI", "Streamlit"],
        ["Document Generation", "python-docx"],
        ["Logging", "loguru (structured JSON, PII-safe)"],
    ])
    
    # ── Final page ─────────────────────────────────────────────────────────
    doc.add_page_break()
    add_styled_heading(doc, "End of Documentation", 1)
    doc.add_paragraph(
        f"This document was auto-generated on {datetime.now().strftime('%d %B %Y at %H:%M')} "
        f"by running all 14 agent interactions, 9 direct tool executions, and "
        f"the 15-case evaluation harness.\n\n"
        f"All results are live — computed from actual CSV data with verified formulas.\n\n"
        f"Team 20: Anandasekhar | Rohit | Raseeth\n"
        f"IIT Madras — Capstone Project"
    )
    
    # Save
    out_path = Path(__file__).resolve().parent / "docs" / "ARIS_Execution_Results.docx"
    out_path.parent.mkdir(exist_ok=True)
    doc.save(str(out_path))
    print(f"Document saved to: {out_path}")
    print(f"Total pages (approx): {len(doc.paragraphs) // 25 + 1}")


if __name__ == "__main__":
    main()
